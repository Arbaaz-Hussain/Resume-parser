import os
import re
import docx
from pdfminer.high_level import extract_text
from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError

def extract_info_from_docx(file_path):
    doc = docx.Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_info_from_pdf(file_path):
    return extract_text(file_path)

def extract_emails(text):
    # Regular expression pattern to match valid email addresses
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    # Find all email addresses in the text using the pattern
    emails = re.findall(email_pattern, text)
    return emails


def extract_phone_numbers(text):
    phone_numbers = re.findall(r'(?:(?:\+?(\d{1,3}))?[-. (]*(\d{3})[-. )]*(\d{3})[-. ]*(\d{4})(?: *x(\d+))?)', text)
    return [''.join(phone) for phone in phone_numbers]

def remove_non_printable(text):
    return ''.join(char for char in text if char.isprintable())

def process_cv(file_path):
    if file_path.endswith('.pdf'):
        text = extract_info_from_pdf(file_path)
    elif file_path.endswith('.docx'):
        text = extract_info_from_docx(file_path)
    else:
        text = ""
    text = remove_non_printable(text)
    emails = extract_emails(text)
    phone_numbers = extract_phone_numbers(text)
    return text, emails, phone_numbers

def process_bundle_cv(folder_path):
    wb = Workbook()
    ws = wb.active
    ws.append(['File Name', 'Email', 'Phone Number', 'Text'])
    for file_name in os.listdir(folder_path):
        if file_name.endswith('.pdf') or file_name.endswith('.docx'):
            file_path = os.path.join(folder_path, file_name)
            try:
                text, emails, phone_numbers = process_cv(file_path)
                ws.append([file_name, ', '.join(emails), ', '.join(phone_numbers), text])
            except IllegalCharacterError:
                print(f"Illegal characters found in {file_name}. Skipping...")
    output_file = os.path.join(folder_path, 'cv_info.xlsx')
    wb.save(output_file)
    return output_file
